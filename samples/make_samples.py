"""Builds a SYNTHETIC production (fictional patient, no real PHI) and a draft Claims Review with
planted problems, reviewer tracked changes and comments.

Planted problems the checker should catch:
  - repeat lactate stated as 1.8 (record says 1.3)                 -> CONTRADICTED
  - urine culture cited to MRX000106 (it is on MRX000107)          -> WRONG_PAGE
  - DRG cited to MRX000111 (not produced)                          -> MISSING_PAGE
  - platelets 98 K/uL stated with no citation (record says 210)    -> UNCITED
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).parent
PAGES = {
    "MRX000101": """EMERGENCY DEPARTMENT TRIAGE NOTE
Patient: Patient A (synthetic)   DOB: 01/15/1961   Arrival: 03/03/2026 21:14
Chief complaint: fever and dysuria x 2 days.
Vital signs: Temp 38.9 C | HR 118 | RR 22 | BP 96/58 | SpO2 95% on room air.
Acuity: ESI 2. Alert and oriented x3.""",
    "MRX000102": """EMERGENCY DEPARTMENT PHYSICIAN NOTE  03/03/2026 21:40
HPI: 65-year-old with 2 days of dysuria and fever. No confusion.
Assessment: urinary tract infection, rule out sepsis.
Plan: blood cultures x2, urine culture, ceftriaxone 1 g IV, lactated Ringer's 30 mL/kg bolus, admit to medicine.""",
    "MRX000103": """LABORATORY RESULTS  Collected 03/03/2026 22:05
WBC 15.2 K/uL (H) | Hemoglobin 13.1 g/dL | Platelets 210 K/uL
Lactate 2.4 mmol/L (H) | Creatinine 1.1 mg/dL | Total bilirubin 0.8 mg/dL
Sodium 136 mEq/L | Potassium 4.0 mEq/L""",
    "MRX000104": """LABORATORY RESULTS  Collected 03/04/2026 06:00
WBC 12.1 K/uL | Platelets 198 K/uL
Lactate 1.3 mmol/L | Creatinine 1.0 mg/dL""",
    "MRX000105": """NURSING FLOWSHEET  03/04/2026
02:00 BP 108/66 HR 102 | 06:00 BP 118/72 HR 94 MAP 87 | 10:00 BP 121/75 HR 88
No vasopressors. Tolerating oral intake. Mentation at baseline.""",
    "MRX000106": """HOSPITALIST HISTORY AND PHYSICAL  03/04/2026 08:15
Assessment: sepsis due to urinary tract infection, present on admission. Hemodynamically stable after fluids.
No acute organ dysfunction identified. Continue ceftriaxone pending culture results.""",
    "MRX000107": """MICROBIOLOGY
Urine culture (03/03/2026): >100,000 CFU/mL Escherichia coli, pan-sensitive.
Blood cultures x2 (03/03/2026): no growth at 5 days.""",
    "MRX000108": """DISCHARGE SUMMARY  Discharged 03/06/2026
Principal diagnosis: Sepsis due to E. coli (A41.51). Secondary: urinary tract infection (N39.0).
Length of stay: 3 days. Discharged home on oral cephalexin.""",
    "MRX000109": """PHYSICIAN PROGRESS NOTE  03/05/2026
Afebrile x 24 hours. WBC trending down. Plan transition to oral antibiotics, anticipate discharge tomorrow.""",
    "MRX000110": """PAYER DENIAL LETTER (excerpt)
Clinical validation review: the documentation does not support sepsis as defined by Sepsis-3.
No organ dysfunction (SOFA increase of 2 or more) is documented. The claim is adjusted to the DRG for urinary tract infection.""",
}

for bates, text in PAGES.items():
    (HERE / "production").mkdir(exist_ok=True)
    (HERE / "production" / f"{bates}.txt").write_text(text + "\n", encoding="utf-8")

doc = Document()
doc.add_heading("Claims Review: Patient A (synthetic demo)", level=0)
doc.add_paragraph("Prepared for counsel. Draft for RN and coder review. All facts must tie to the produced record.")

doc.add_heading("Summary of Records", level=1)
summary = [
    "Patient A presented to the emergency department on 03/03/2026 with fever and dysuria (MRX000101).",
    "Triage vital signs were temperature 38.9 °C, heart rate 118, respiratory rate 22 and blood pressure 96/58 (MRX000101).",
    "Initial labs showed a white blood cell count of 15.2 K/uL and lactate of 2.4 mmol/L (MRX000103).",
    "Repeat lactate on 03/04/2026 was 1.8 mmol/L (MRX000104).",
    "Urine culture grew >100,000 CFU/mL of E. coli and blood cultures showed no growth at 5 days (MRX000106).",
    "The patient received ceftriaxone 1 g IV and a 30 mL/kg fluid bolus (MRX000102).",
    "No vasopressors were required and blood pressure improved to 118/72 (MRX000105).",
    "The patient was discharged on 03/06/2026 with a principal diagnosis of sepsis, A41.51 (MRX000108).",
    "The claim was billed as MS-DRG 872 (MRX000111).",
]
paras = [doc.add_paragraph(s) for s in summary]

doc.add_heading("Analysis", level=1)
analysis = [
    "Platelets were 98 K/uL, indicating thrombocytopenia.",
    "The hospitalist documented that no acute organ dysfunction was identified (MRX000106).",
    "Under Sepsis-3, sepsis requires an acute increase in SOFA score of two or more points due to infection, which the record does not document.",
]
apar = [doc.add_paragraph(s) for s in analysis]


def tracked(paragraph, kind: str, text: str, author: str, date: str):
    """Append a tracked insertion or deletion run (python-docx has no API for this; write the XML)."""
    wrapper = OxmlElement("w:ins" if kind == "ins" else "w:del")
    wrapper.set(qn("w:id"), str(abs(hash((text, kind))) % 10_000))
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), date)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t" if kind == "ins" else "w:delText")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    wrapper.append(run)
    paragraph._p.append(wrapper)


RN = "J. Mathew, RN"
tracked(paras[5], "ins", " Antibiotics were started within 1 hour of triage.", RN, "2026-10-01T10:12:00Z")
tracked(apar[2], "del", " clearly", RN, "2026-10-01T10:20:00Z")
doc.add_comment([r for r in paras[3].runs], text="Please confirm the repeat lactate value against the 03/04 lab page.", author=RN, initials="JM")
doc.add_comment([r for r in paras[8].runs], text="DRG 872 title should read 'without MCC'; also confirm the page cite.", author="S. Rao, CCS", initials="SR")

doc.save(HERE / "draft_claims_review.docx")
print(f"wrote {len(PAGES)} pages to samples/production and samples/draft_claims_review.docx")
