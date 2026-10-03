"""Read the draft report (.docx): cited sentences, reviewer tracked changes, and comments."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn

from .records import find_cites

W_INS, W_DEL, W_T, W_DELTEXT = qn("w:ins"), qn("w:del"), qn("w:t"), qn("w:delText")


@dataclass
class Sentence:
    section: str
    paragraph_index: int
    text: str
    cites: list[str] = field(default_factory=list)


@dataclass
class ReviewItem:
    kind: str           # insertion | deletion | comment
    author: str
    date: str
    text: str
    anchor: str         # for comments: the highlighted text; for edits: the paragraph
    section: str


def paragraph_text(p_el, include_deleted: bool = False) -> str:
    """Text of a paragraph as currently shown with tracked changes accepted (or with deletions kept)."""
    parts = []
    for node in p_el.iter():
        if node.tag == W_T:
            parts.append(node.text or "")
        elif node.tag == W_DELTEXT and include_deleted:
            parts.append(node.text or "")
    return "".join(parts)


def _sentences(text: str) -> list[str]:
    # Split after a period that ends a sentence, keeping citations like "(MRX000101)." attached.
    chunks = re.split(r"(?<=[.!?])\s+(?=[A-Z(])", text.strip())
    return [c.strip() for c in chunks if c.strip()]


def read_report(path) -> tuple[list[Sentence], list[ReviewItem], object]:
    doc = Document(str(path))
    sentences: list[Sentence] = []
    section = "Front matter"
    sections_by_par: list[str] = []
    for i, p in enumerate(doc.paragraphs):
        if p.style is not None and p.style.name.lower().startswith("heading"):
            section = p.text.strip() or section
        sections_by_par.append(section)
        text = paragraph_text(p._p)
        for s in _sentences(text):
            sentences.append(Sentence(section, i, s, find_cites(s)))

    review: list[ReviewItem] = []
    paragraphs = doc.paragraphs  # keep the proxies alive: lxml element identity is only stable while referenced
    par_index = {p._p: i for i, p in enumerate(paragraphs)}
    for tag, kind, text_tag in ((W_INS, "insertion", W_T), (W_DEL, "deletion", W_DELTEXT)):
        for el in doc.element.body.iter(tag):
            p_el = el.getparent()
            while p_el is not None and p_el.tag != qn("w:p"):
                p_el = p_el.getparent()
            idx = par_index.get(p_el, 0) if p_el is not None else 0
            review.append(ReviewItem(kind, el.get(qn("w:author"), ""), el.get(qn("w:date"), "")[:10],
                                     "".join(t.text or "" for t in el.iter(text_tag)),
                                     paragraph_text(p_el, include_deleted=True) if p_el is not None else "",
                                     sections_by_par[idx] if idx < len(sections_by_par) else ""))

    anchors = _comment_anchors(doc)
    for c in doc.comments:
        anchor, idx = anchors.get(str(c.comment_id), ("", 0))
        review.append(ReviewItem("comment", c.author or "", c.timestamp.date().isoformat() if c.timestamp else "",
                                 c.text, anchor, sections_by_par[idx] if idx < len(sections_by_par) else ""))
    return sentences, review, doc


def _comment_anchors(doc) -> dict[str, tuple[str, int]]:
    """Text between each commentRangeStart/End pair, and the paragraph it starts in."""
    out: dict[str, tuple[str, int]] = {}
    open_ranges: dict[str, list[str]] = {}
    start_par: dict[str, int] = {}
    for i, p in enumerate(doc.paragraphs):
        for node in p._p.iter():
            if node.tag == qn("w:commentRangeStart"):
                cid = node.get(qn("w:id"))
                open_ranges[cid] = []
                start_par[cid] = i
            elif node.tag == qn("w:commentRangeEnd"):
                cid = node.get(qn("w:id"))
                out[cid] = ("".join(open_ranges.pop(cid, [])), start_par.get(cid, i))
            elif node.tag == W_T:
                for parts in open_ranges.values():
                    parts.append(node.text or "")
    return out
