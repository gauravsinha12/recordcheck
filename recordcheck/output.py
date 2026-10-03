"""Excel workbook for the team, and an annotated copy of the draft with a Word comment per issue."""
from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .verify import ORDER, Check

FILL = {
    "MISSING_PAGE": "F8D7DA", "CONTRADICTED": "F8D7DA", "WRONG_PAGE": "FFF3CD",
    "UNCITED": "FFF3CD", "SUPPORTED": "D1E7DD", "NOT_CHECKED": "EEEEEE",
}


def write_workbook(path, checks: list[Check], review, report_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    counts = Counter(c.status for c in checks)
    ws.append(["Provenance check", report_name])
    ws.append([])
    ws.append(["Status", "Sentences", "Meaning"])
    meaning = {
        "MISSING_PAGE": "Cited page is not in the production",
        "CONTRADICTED": "A stated value is not on the cited page or anywhere in the record",
        "WRONG_PAGE": "Value exists in the record, but on a different page",
        "UNCITED": "Clinical values stated with no citation",
        "SUPPORTED": "Every value found on the cited page(s)",
        "NOT_CHECKED": "No machine-checkable values; reviewer judgment",
    }
    for status in ORDER:
        ws.append([status, counts.get(status, 0), meaning[status]])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=FILL[status])
    ws.append([])
    ws.append(["Reviewer edits and comments to work back in", len(review)])
    ws["A1"].font = Font(bold=True, size=13)
    _widths(ws, [24, 18, 70])

    cs = wb.create_sheet("Citation check")
    cs.append(["Status", "Section", "Sentence", "Cited pages", "Values checked", "Not on cited page", "Found on", "Note", "LLM verdict", "LLM quote / issue"])
    for c in sorted(checks, key=lambda c: (ORDER.index(c.status), c.sentence.paragraph_index)):
        found = "; ".join(f"{k} → {', '.join(v)}" for k, v in c.found_elsewhere.items() if v)
        cs.append([c.status, c.sentence.section, c.sentence.text, ", ".join(c.sentence.cites),
                   ", ".join(f.text for f in c.facts), ", ".join(f.text for f in c.missing), found, c.note,
                   (c.llm or {}).get("verdict", ""), ((c.llm or {}).get("quote") or (c.llm or {}).get("issue") or "")])
        cs.cell(cs.max_row, 1).fill = PatternFill("solid", fgColor=FILL[c.status])
    _header(cs)
    _widths(cs, [15, 20, 70, 22, 30, 22, 26, 50, 14, 50])

    rs = wb.create_sheet("Reviewer edits")
    rs.append(["Done?", "Type", "Author", "Date", "Section", "Text", "Context / highlighted text"])
    for r in review:
        rs.append(["", r.kind, r.author, r.date, r.section, r.text, r.anchor])
    _header(rs)
    _widths(rs, [8, 11, 22, 11, 20, 50, 70])
    wb.save(path)


def _header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def annotate(doc, checks: list[Check], out_path) -> int:
    """Add a Word comment on every paragraph with an issue, authored 'Provenance Check'."""
    flagged = 0
    paragraphs = doc.paragraphs
    for c in checks:
        if c.status in ("SUPPORTED", "NOT_CHECKED"):
            continue
        p = paragraphs[c.sentence.paragraph_index]
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            continue
        doc.add_comment(runs, text=f"{c.status}: {c.note}\nSentence: “{c.sentence.text[:160]}”", author="Provenance Check", initials="PC")
        flagged += 1
    doc.save(str(out_path))
    return flagged
